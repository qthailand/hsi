"""Spectra export of the cube viewer: plain CSV, XLSX workbook and clipboard TSV.

All three share one table built by :func:`spectra_table`::

    wavelength,"(120,48)_mean","(120,48)_std",leaf-mean-0_mean,leaf-mean-0_std
    854.17,812.0,0.0,1503.2215338,41.0634577

* First column ``wavelength`` (nm, the cube's own band order - never sorted, so a
  descending cube stays descending) or ``band`` (0-based index) when the cube has no
  wavelengths or their count differs from the band count (the plot does the same).
* Then ``<name>_mean`` and ``<name>_std`` for every exported record, in the order given
  (creation order from ``SpectrumSet``).  A point is a single pixel: its std is 0.0.
  Header cells holding a comma or a quote (point names ``(ix,iy)``) are quoted by ``csv``.
* Values: float64 statistics are written with ``repr`` (shortest text that round-trips
  bit-exactly).  float32 wavelengths are written as the shortest decimal that round-trips
  to the stored float32 (``854.17``, not ``854.1699829101562``).  NaN (a band with no
  valid pixel) is an empty cell; +-inf is ``inf``/``-inf`` in CSV/TSV and empty in XLSX.

Which records are exported: those with ``state == 'ready'`` and stats, and - unless
``visible_only=False`` - ``visible``.  Pending and failed records are always skipped.
Names must be unique.  Every writer raises ``ValueError`` when nothing is exportable.
``show_std`` only affects the plot; exports always carry the std column.

CSV (:func:`export_csv`): a plain table exactly in the style of the user's
``PVC_reflectance_with_std.csv`` - header row, then numbers; no comment or metadata rows;
CRLF line ends (``csv`` default); UTF-8 without BOM, or ``utf-8-sig`` when a header has
non-ASCII text so Excel shows Thai names.  The CSV deliberately carries no
``format_version`` (user decision 2026-10-19, matching the PVC files; an exception to
MODULE_CONTRACT "File formats"); the XLSX ``Metadata`` sheet carries ``format_version=1``
and every other metadata item.  To read the numbers back bit-exactly use ``csv`` plus
``float()`` or ``pandas.read_csv(path, float_precision="round_trip")``; plain
``pandas.read_csv(path)`` uses pandas' faster default float parser, which is only
accurate to about 1e-12 relative (pandas 2.2: about 40 % of ``repr`` strings come back a
few ulp off, e.g. ``0.00010856208701859327`` -> ``0.0001085620870185``).

XLSX (:func:`export_xlsx`, needs ``openpyxl``; ``HAVE_OPENPYXL`` tells): sheets exactly

* ``Spectra``  - the same table (panes frozen at B2).
* ``ROIs``     - one row per exported record: rid, name, class, kind, n_px, min_n_valid,
  max_n_saturated, pixel bbox c0, r0, c1, r1 (half-open; for a lasso the pixel-centre
  bbox of its vertices, clipped when ``metadata['shape']`` is given), n_vertices,
  vertices_json (blank with a note above Excel's 32,767-character cell limit), color,
  visible, created_at, method, note.
* ``Metadata`` - key/value rows: ``format_version`` = 1, exported_at, std_ddof = 0, the
  pixel rule, wavelength order, ..., then the caller's ``metadata`` (``source``,
  ``units``, ``dtype``, ``shape``, ``axis_order`` always appear; camera scalars such as
  ``m``, ``c``, ``pixel_size``, ``camera_bin``, ``bit_depth``, ``full_scale`` follow).

openpyxl formats floats with ``%.16g`` (not round-trip exact), so numbers are written
as numeric cells holding ``repr`` text - openpyxl and ``pandas.read_excel`` read them
back bit-exactly (tested); integral floats stay floats (``400.0``, not ``400``).

Clipboard (:func:`spectra_text`): the same table as TSV, ``\\n`` line ends, no comment.

Files are written to a temporary file in the target folder and moved into place, so a
failed export (e.g. the target is open in Excel) never leaves a half-written file.

Data conventions:
    One row per band in the cube's own order; ``wavelengths`` is a 1-D array in nm (one
    per band) or None.  Values are the records' float64 per-band statistics in the cube's
    units (counts for integer cubes), std with ddof=0; the records and their arrays are
    only read.  Lasso vertices in the ROIs sheet are image-corner ``[x, y]`` floats (x =
    column = scan, y = row = slit), bboxes half-open pixel boxes.

Threading:
    Pure functions; safe in any thread.  ~ms for typical tables (see README).

Example:
    >>> from modules.cube_viewer.spectra_export import export_csv     # doctest: +SKIP
    >>> export_csv("leaf_spectra.csv", spectrum_set.records(), wavelengths)  # doctest: +SKIP
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import json
import math
import os
import uuid

import numpy as np

from .roi_stats import snap_rect

try:  # optional dependency: XLSX export only
    import openpyxl
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover - depends on the environment
    openpyxl = None
    HAVE_OPENPYXL = False

__all__ = ["FORMAT_VERSION", "HAVE_OPENPYXL", "PIXEL_RULE", "STD_DDOF", "SHEETS",
           "ROI_COLUMNS", "exportable", "spectra_table", "export_csv", "spectra_text",
           "export_xlsx"]

FORMAT_VERSION = 1
STD_DDOF = 0
SHEETS = ("Spectra", "ROIs", "Metadata")
PIXEL_RULE = ("pixel (c, r) spans [c, c+1) x [r, r+1); rect: pixels whose centre lies in "
              "the closed dragged rectangle (c0 = ceil(xmin - 0.5), c1 = floor(xmax - 0.5) + 1); "
              "lasso: pixel centres inside the polygon, even-odd rule; point: the clicked pixel")
ROI_COLUMNS = ("rid", "name", "class", "kind", "n_px", "min_n_valid", "max_n_saturated",
               "c0", "r0", "c1", "r1", "n_vertices", "vertices_json", "color", "visible",
               "created_at", "method", "note")

_EXCEL_CELL_MAX = 32767
_RECOMMENDED_META = ("source", "units", "dtype", "shape", "axis_order")
_DEFAULT_AXIS_ORDER = "(slit, scan, wavelength) = (rows, columns, bands)"
_BIG = 1 << 40   # stands in for an unknown image size when clipping a lasso bbox


# ── shared table ────────────────────────────────────────────

def exportable(records, *, visible_only: bool = True) -> list:
    """Records the writers export: ready with stats, and visible unless ``visible_only=False``."""
    return [r for r in records
            if r.state == "ready" and r.stats is not None and (r.visible or not visible_only)]


def _axis(wavelengths, n_bands: int) -> tuple:
    """``('wavelength', values)`` when usable, else ``('band', 0..n_bands-1)``."""
    if wavelengths is not None:
        wl = np.asarray(wavelengths)
        if wl.ndim == 1 and wl.shape[0] == n_bands and wl.dtype.kind in "fiu":
            if wl.dtype.kind == "f" and wl.dtype.itemsize < 8:
                # shortest decimal that round-trips to the stored float32/float16
                return "wavelength", [float(str(v)) for v in wl]
            if wl.dtype.kind == "f":
                return "wavelength", wl.astype(np.float64).tolist()
            return "wavelength", wl.tolist()
    return "band", list(range(n_bands))


def spectra_table(records, wavelengths, *, visible_only: bool = True) -> tuple:
    """Build the export table shared by CSV, XLSX and the clipboard.

    Args:
        records: Iterable of ``SpectrumRecord`` (see :func:`exportable` for the filter).
        wavelengths: 1-D wavelengths in nm (one per band) or ``None``.
        visible_only: Skip hidden records (default).

    Returns:
        tuple: ``(header, rows)`` - ``header`` is a list of str (``'wavelength'`` or
        ``'band'``, then ``<name>_mean``, ``<name>_std`` per record); ``rows`` is one list
        per band in the cube's band order holding the axis value (float or int) and
        Python floats (NaN = missing).

    Raises:
        ValueError: Nothing exportable, duplicate names, or records with different band
            counts.
    """
    recs = exportable(records, visible_only=visible_only)
    if not recs:
        raise ValueError("no spectra to export: none is ready"
                         + (" and visible" if visible_only else ""))
    n_bands = int(np.shape(recs[0].stats.mean)[0])
    seen = set()
    header_names = []
    columns = []
    for record in recs:
        if record.name in seen:
            raise ValueError(f"duplicate spectrum name {record.name!r}")
        seen.add(record.name)
        mean = np.asarray(record.stats.mean, dtype=np.float64)
        std = np.asarray(record.stats.std, dtype=np.float64)
        if mean.shape != (n_bands,) or std.shape != (n_bands,):
            raise ValueError(f"spectrum {record.name!r} has {mean.shape} bands, expected ({n_bands},)")
        header_names += [f"{record.name}_mean", f"{record.name}_std"]
        columns += [mean.tolist(), std.tolist()]
    axis_name, axis_values = _axis(wavelengths, n_bands)
    rows = [list(row) for row in zip(axis_values, *columns)]
    return [axis_name] + header_names, rows


def _text(value) -> str:
    """CSV/TSV text of one table value: ``repr`` for floats, empty for NaN/None."""
    if value is None:
        return ""
    if isinstance(value, float):
        return "" if value != value else repr(value)
    return str(value)


def _text_rows(header, rows) -> list:
    return [header] + [[_text(v) for v in row] for row in rows]


# ── file helpers ────────────────────────────────────────────

def _atomic_write(path: str, write) -> None:
    """Call ``write(tmp_path)`` then move the temporary file onto ``path``.

    The temporary file sits next to ``path`` (same volume, so the move is atomic) under a
    random name that ``write`` creates itself, so the result gets the normal file
    permissions (``mkstemp`` would leave it owner-only on POSIX).
    """
    folder, base = os.path.split(os.path.abspath(path))
    tmp = os.path.join(folder, f".{base}.{uuid.uuid4().hex[:12]}.tmp")
    try:
        write(tmp)
        os.replace(tmp, path)
    except BaseException:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise


# ── CSV / TSV ───────────────────────────────────────────────

def export_csv(path, records, wavelengths, *, visible_only: bool = True) -> str:
    """Write the spectra table as a plain PVC-style CSV (see the module docstring).

    Args:
        path: Target file; ``'.csv'`` is appended only when it has no extension.
        records: ``SpectrumRecord`` iterable.
        wavelengths: 1-D wavelengths in nm or ``None`` (-> ``band`` column).
        visible_only: Skip hidden records (default).

    Returns:
        str: The path written.

    Raises:
        ValueError: See :func:`spectra_table`.
        OSError: The file cannot be written (e.g. open in Excel); the target is unchanged.
    """
    header, rows = spectra_table(records, wavelengths, visible_only=visible_only)
    path = os.fspath(path)
    if not os.path.splitext(path)[1]:
        path += ".csv"
    encoding = "utf-8" if all(h.isascii() for h in header) else "utf-8-sig"
    lines = _text_rows(header, rows)

    def write(tmp):
        with open(tmp, "w", encoding=encoding, newline="") as handle:
            csv.writer(handle).writerows(lines)

    _atomic_write(path, write)
    return path


def spectra_text(records, wavelengths, *, visible_only: bool = True) -> str:
    """The spectra table as tab-separated text for the clipboard (``\\n`` line ends).

    Raises:
        ValueError: See :func:`spectra_table`.
    """
    header, rows = spectra_table(records, wavelengths, visible_only=visible_only)
    buffer = io.StringIO()
    csv.writer(buffer, delimiter="\t", lineterminator="\n").writerows(_text_rows(header, rows))
    return buffer.getvalue()


# ── XLSX ────────────────────────────────────────────────────

def _xl_text(ws, value: str, bold: bool = False):
    """Text cell that never turns into a formula and holds no illegal XML characters."""
    cell = WriteOnlyCell(ws, value=ILLEGAL_CHARACTERS_RE.sub("", value))
    cell.data_type = "s"
    if bold:
        cell.font = Font(bold=True)
    return cell


def _xl_value(ws, value):
    """Cell for one value; floats keep every bit (numeric cell holding ``repr`` text)."""
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        value = int(value)
        if abs(value) < 10 ** 15:
            return value
        value = float(value)
    if isinstance(value, (float, np.floating)):
        value = float(value)
        if not math.isfinite(value):
            return None
        cell = WriteOnlyCell(ws, value=repr(value))
        cell.data_type = "n"
        return cell
    if isinstance(value, str):
        return _xl_text(ws, value)
    if isinstance(value, np.ndarray):
        value = value.tolist()
    if isinstance(value, (list, tuple, dict)):
        return _xl_text(ws, json.dumps(value, ensure_ascii=False, default=_json_default))
    return _xl_text(ws, str(value))


def _json_default(value):
    """JSON fallback for numpy scalars/arrays and other objects."""
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    return str(value)


def _set_widths(ws, widths) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width))


def _image_hw(metadata: dict):
    """``(H, W)`` from ``metadata['shape']`` when it is a usable sequence, else ``None``."""
    shape = metadata.get("shape")
    try:
        height, width = int(shape[0]), int(shape[1])
    except (TypeError, ValueError, IndexError, KeyError):
        return None
    return (height, width) if height > 0 and width > 0 else None


def _record_bbox(record, image_hw) -> tuple:
    """Half-open pixel bbox ``(c0, r0, c1, r1)`` of a record."""
    geom = record.geometry
    if all(k in geom for k in ("c0", "r0", "c1", "r1")):
        return tuple(int(geom[k]) for k in ("c0", "r0", "c1", "r1"))
    if record.kind == "point":
        x, y = int(geom["x"]), int(geom["y"])
        return x, y, x + 1, y + 1
    vertices = np.asarray(geom["vertices"], dtype=np.float64)
    height, width = image_hw if image_hw else (_BIG, _BIG)
    return tuple(int(v) for v in snap_rect(float(vertices[:, 0].min()), float(vertices[:, 1].min()),
                                           float(vertices[:, 0].max()), float(vertices[:, 1].max()),
                                           height, width))


def _roi_row(ws, record, image_hw) -> list:
    stats = record.stats
    c0, r0, c1, r1 = _record_bbox(record, image_hw)
    vertices = record.geometry.get("vertices") if record.kind in ("lasso", "polygon") else None
    n_vertices = len(vertices) if vertices is not None else None
    vertices_json = None
    note = ""
    if vertices is not None:
        vertices_json = json.dumps(vertices, separators=(",", ":"), default=_json_default)
        if len(vertices_json) > _EXCEL_CELL_MAX:
            note = (f"vertices omitted: {len(vertices_json):,} characters exceed the Excel "
                    f"cell limit ({_EXCEL_CELL_MAX:,})")
            vertices_json = None
    n_valid = np.asarray(stats.n_valid)
    n_sat = np.asarray(stats.n_saturated)
    values = [record.rid, record.name, record.class_name, record.kind, int(stats.n_px),
              int(n_valid.min()) if n_valid.size else 0,
              int(n_sat.max()) if n_sat.size else 0,
              c0, r0, c1, r1, n_vertices, vertices_json, record.color, bool(record.visible),
              record.created_at, stats.method, note]
    return [_xl_value(ws, v) for v in values]


def _wavelength_order(axis_name: str, axis_values) -> str:
    if axis_name != "wavelength":
        return "n/a (band index)"
    diffs = np.diff(np.asarray(axis_values, dtype=np.float64))
    if diffs.size == 0 or np.all(diffs > 0):
        return "ascending"
    if np.all(diffs < 0):
        return "descending"
    return "unordered (cube band order)"


def _metadata_rows(header, rows, n_spectra: int, metadata: dict) -> list:
    axis_name = header[0]
    fixed = [
        ("format_version", FORMAT_VERSION),
        ("exported_at", _dt.datetime.now().astimezone().isoformat(timespec="seconds")),
        ("writer", "modules.cube_viewer.spectra_export"),
        ("table_layout", f"{axis_name}, then <name>_mean, <name>_std per spectrum; "
                         "rows in the cube's band order; empty cell = no valid pixel"),
        ("axis", axis_name),
        ("wavelength_units", "nm" if axis_name == "wavelength" else ""),
        ("wavelength_order", _wavelength_order(axis_name, [r[0] for r in rows])),
        ("n_bands", len(rows)),
        ("n_spectra", n_spectra),
        ("std_ddof", STD_DDOF),
        ("std_definition", "population standard deviation over the valid pixels of each band "
                           "(ddof=0); a single-pixel point has std 0"),
        ("pixel_rule", PIXEL_RULE),
        ("csv_note", "the CSV export is a plain table without format_version or metadata "
                     "(user decision 2026-10-19); this sheet carries them"),
    ]
    owned = {key for key, _ in fixed}
    out = list(fixed)
    for key in _RECOMMENDED_META:
        default = _DEFAULT_AXIS_ORDER if key == "axis_order" else None
        out.append((key, metadata.get(key, default)))
    taken = owned | set(_RECOMMENDED_META)
    out += [(str(k), v) for k, v in metadata.items() if str(k) not in taken]
    return out


def export_xlsx(path, records, wavelengths, metadata=None, *, visible_only: bool = True) -> str:
    """Write the spectra workbook (sheets ``Spectra``, ``ROIs``, ``Metadata``).

    Args:
        path: Target file; ``'.xlsx'`` is appended when it does not end with it.
        records: ``SpectrumRecord`` iterable.
        wavelengths: 1-D wavelengths in nm or ``None`` (-> ``band`` column).
        metadata: Extra key/value items for the ``Metadata`` sheet, e.g. ``source``,
            ``units``, ``dtype``, ``shape`` (H, W, B - also clips lasso bboxes),
            ``axis_order`` and camera scalars.  Keys the writer owns
            (``format_version``, ``exported_at``, ``std_ddof``, ...) cannot be overridden.
        visible_only: Skip hidden records (default).

    Returns:
        str: The path written.

    Raises:
        ImportError: openpyxl is not installed.
        ValueError: See :func:`spectra_table`.
        OSError: The file cannot be written; the target is unchanged.
    """
    if not HAVE_OPENPYXL:
        raise ImportError("XLSX export needs openpyxl: pip install openpyxl")
    metadata = dict(metadata or {})
    records = list(records)
    header, rows = spectra_table(records, wavelengths, visible_only=visible_only)
    recs = exportable(records, visible_only=visible_only)
    path = os.fspath(path)
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet(SHEETS[0])
    ws.freeze_panes = "B2"
    _set_widths(ws, [len(h) + 2 for h in header])
    ws.append([_xl_text(ws, h, bold=True) for h in header])
    for row in rows:
        ws.append([_xl_value(ws, v) for v in row])

    ws = wb.create_sheet(SHEETS[1])
    ws.freeze_panes = "C2"
    _set_widths(ws, [len(c) + 2 for c in ROI_COLUMNS])
    ws.append([_xl_text(ws, c, bold=True) for c in ROI_COLUMNS])
    image_hw = _image_hw(metadata)
    for record in recs:
        ws.append(_roi_row(ws, record, image_hw))

    ws = wb.create_sheet(SHEETS[2])
    _set_widths(ws, [24, 60])
    ws.append([_xl_text(ws, "key", bold=True), _xl_text(ws, "value", bold=True)])
    for key, value in _metadata_rows(header, rows, len(recs), metadata):
        ws.append([_xl_text(ws, key), _xl_value(ws, value)])

    _atomic_write(path, wb.save)
    return path
